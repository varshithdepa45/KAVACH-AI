"""Deliverable generation service for KAVACH AI.

Always writes a .txt and a .md report. Additionally writes .docx (python-docx),
.xlsx (openpyxl) and .pdf (reportlab or fpdf2) ONLY when the corresponding
library is importable. All optional imports are guarded; failure degrades
gracefully to the always-available formats.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .. import config
from . import scenario

# --- Optional deps (guarded) --------------------------------------------------
try:
    import docx  # type: ignore
    _DOCX = True
except Exception:  # noqa: BLE001
    _DOCX = False

try:
    import openpyxl  # type: ignore
    _XLSX = True
except Exception:  # noqa: BLE001
    _XLSX = False

_PDF_LIB = None
try:
    from reportlab.lib.pagesizes import letter  # type: ignore
    from reportlab.pdfgen import canvas as _rl_canvas  # type: ignore
    _PDF_LIB = "reportlab"
except Exception:  # noqa: BLE001
    try:
        from fpdf import FPDF  # type: ignore
        _PDF_LIB = "fpdf2"
    except Exception:  # noqa: BLE001
        _PDF_LIB = None


def optional_formats() -> dict[str, Any]:
    return {"docx": _DOCX, "xlsx": _XLSX, "pdf": bool(_PDF_LIB), "pdf_lib": _PDF_LIB}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _report_lines(ctx) -> list[str]:
    v = ctx.artifacts.get("verification", {})
    lines = [
        f"{config.APP_NAME} - Autonomous Inspection Report",
        f"Facility: {scenario.FACILITY} | {scenario.UNIT}",
        f"Generated: {_now()} | Mode: {config.KAVACH_MODE}",
        f"Run ID: {ctx.artifacts.get('run_id')}",
        "",
        "FICTIONAL DEMONSTRATION DATA - NOT REAL FACILITY DATA",
        "=" * 64,
        "",
        "1. EXECUTIVE SUMMARY",
        f"   Verification score: {ctx.verification_score}%",
        f"   Evidence-backed: {ctx.evidence_backed}",
        f"   Model routing: {ctx.routing.get('model')} ({ctx.routing.get('task_type')})",
        "",
        "2. FINDINGS",
    ]
    for i, f in enumerate(ctx.findings, start=1):
        flag = " [NEEDS HUMAN REVIEW]" if f.get("needs_review") else ""
        lines.append(f"   {i}. [{f['severity'].upper()}] {f['title']}{flag}")
        lines.append(f"      Equipment: {f['equipment_id']} | Confidence: {f['confidence']}%")
        lines.append(f"      {f['description']}")
        for ev in f.get("evidence", []):
            lines.append(f"      - Evidence: {ev['source']} (p.{ev.get('page')}) "
                         f"[{ev.get('confidence')}%]")
            lines.append(f"        \"{ev.get('excerpt', '')}\"")
        lines.append("")
    lines += [
        "3. VERIFICATION",
        f"   Score: {v.get('score')}% | Evidence-backed: {v.get('evidence_backed')}",
        f"   Findings flagged for review: {', '.join(v.get('needs_human_review', [])) or 'none'}",
        "",
        "4. PIPELINE TRACE",
    ]
    for s in ctx.steps:
        lines.append(f"   - [{s.agent}] {s.message}")
    return lines


def _md_report(ctx) -> str:
    v = ctx.artifacts.get("verification", {})
    md = [
        f"# {config.APP_NAME} - Autonomous Inspection Report",
        "",
        f"**Facility:** {scenario.FACILITY} - {scenario.UNIT}  ",
        f"**Generated:** {_now()}  ",
        f"**Mode:** `{config.KAVACH_MODE}`  |  **Run ID:** {ctx.artifacts.get('run_id')}",
        "",
        "> FICTIONAL DEMONSTRATION DATA - NOT REAL FACILITY DATA",
        "",
        "## Executive Summary",
        "",
        f"- **Verification score:** {ctx.verification_score}%",
        f"- **Evidence-backed:** {ctx.evidence_backed}",
        f"- **Model routing:** {ctx.routing.get('model')} ({ctx.routing.get('task_type')})",
        "",
        "## Findings",
        "",
    ]
    for i, f in enumerate(ctx.findings, start=1):
        flag = " _(needs human review)_" if f.get("needs_review") else ""
        md.append(f"### {i}. {f['title']}{flag}")
        md.append("")
        md.append(f"- **Severity:** {f['severity']}")
        md.append(f"- **Equipment:** {f['equipment_id']}")
        md.append(f"- **Confidence:** {f['confidence']}%")
        md.append(f"- {f['description']}")
        md.append("")
        md.append("| Source | Page | Confidence | Excerpt |")
        md.append("|---|---|---|---|")
        for ev in f.get("evidence", []):
            excerpt = ev.get("excerpt", "").replace("|", "/")
            md.append(f"| {ev['source']} | {ev.get('page')} | {ev.get('confidence')}% | {excerpt} |")
        md.append("")
    md += [
        "## Verification",
        "",
        f"- Score: **{v.get('score')}%**",
        f"- Evidence-backed: **{v.get('evidence_backed')}**",
        f"- Flagged for review: {', '.join(v.get('needs_human_review', [])) or 'none'}",
        "",
        "## Pipeline Trace",
        "",
    ]
    for s in ctx.steps:
        md.append(f"1. **{s.agent}** - {s.message}")
    return "\n".join(md)


def generate_all(ctx) -> list[dict]:
    """Write all deliverables for the run; return list of metadata dicts."""
    config.ensure_dirs()
    run_id = ctx.artifacts.get("run_id", 0)
    base = f"kavach_report_run{run_id}"
    out_dir = config.GENERATED_DIR
    produced: list[dict] = []

    def record(path: Path, fmt: str, name: str) -> None:
        produced.append({
            "name": name,
            "filename": path.name,
            "path": str(path),
            "fmt": fmt,
            "size_bytes": path.stat().st_size if path.exists() else 0,
        })

    lines = _report_lines(ctx)

    # Always: .txt
    txt_path = out_dir / f"{base}.txt"
    txt_path.write_text("\n".join(lines), encoding="utf-8")
    record(txt_path, "txt", "Inspection Report (Text)")

    # Always: .md
    md_path = out_dir / f"{base}.md"
    md_path.write_text(_md_report(ctx), encoding="utf-8")
    record(md_path, "md", "Inspection Report (Markdown)")

    # Optional: .docx
    if _DOCX:
        try:
            doc = docx.Document()
            doc.add_heading(f"{config.APP_NAME} - Inspection Report", level=0)
            doc.add_paragraph(f"{scenario.FACILITY} - {scenario.UNIT}")
            doc.add_paragraph("FICTIONAL DEMONSTRATION DATA - NOT REAL FACILITY DATA")
            doc.add_heading("Findings", level=1)
            for f in ctx.findings:
                doc.add_heading(f["title"], level=2)
                doc.add_paragraph(
                    f"Severity: {f['severity']} | Equipment: {f['equipment_id']} | "
                    f"Confidence: {f['confidence']}%")
                doc.add_paragraph(f["description"])
                for ev in f.get("evidence", []):
                    doc.add_paragraph(
                        f"Evidence: {ev['source']} (p.{ev.get('page')}) - {ev.get('excerpt','')}",
                        style="List Bullet")
            docx_path = out_dir / f"{base}.docx"
            doc.save(str(docx_path))
            record(docx_path, "docx", "Inspection Report (Word)")
        except Exception:  # noqa: BLE001
            pass

    # Optional: .xlsx (findings matrix)
    if _XLSX:
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Findings"
            ws.append(["#", "Title", "Severity", "Equipment", "Confidence%",
                       "Needs Review", "Evidence Count"])
            for i, f in enumerate(ctx.findings, start=1):
                ws.append([i, f["title"], f["severity"], f["equipment_id"],
                           f["confidence"], bool(f["needs_review"]),
                           len(f.get("evidence", []))])
            xlsx_path = out_dir / f"{base}.xlsx"
            wb.save(str(xlsx_path))
            record(xlsx_path, "xlsx", "Findings Matrix (Excel)")
        except Exception:  # noqa: BLE001
            pass

    # Optional: .pdf
    if _PDF_LIB == "reportlab":
        try:
            pdf_path = out_dir / f"{base}.pdf"
            c = _rl_canvas.Canvas(str(pdf_path), pagesize=letter)
            width, height = letter
            y = height - 50
            for line in lines:
                if y < 50:
                    c.showPage()
                    y = height - 50
                c.drawString(40, y, line[:110])
                y -= 14
            c.save()
            record(pdf_path, "pdf", "Inspection Report (PDF)")
        except Exception:  # noqa: BLE001
            pass
    elif _PDF_LIB == "fpdf2":
        try:
            pdf_path = out_dir / f"{base}.pdf"
            pdf = FPDF()
            pdf.set_margins(15, 15, 15)
            pdf.add_page()
            pdf.set_font("Helvetica", size=9)
            for line in lines:
                safe = line.encode("latin-1", "replace").decode("latin-1") or " "
                pdf.multi_cell(0, 5, safe, new_x="LMARGIN", new_y="NEXT")
            pdf.output(str(pdf_path))
            record(pdf_path, "pdf", "Inspection Report (PDF)")
        except Exception:  # noqa: BLE001
            pass

    return produced
